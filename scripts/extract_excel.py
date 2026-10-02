#!/usr/bin/env python3
"""Excel del modelo → cifras de project.json (CLAUDE.md §4). Primera versión.

Uso:
    python3 scripts/extract_excel.py <carpeta-o-slug> --excel RUTA.xlsx [--escribir] [--json SALIDA.json]

- Lee el Excel (que vive FUERA del repo: se descarga de Drive a una carpeta temporal; nunca se versiona).
- Busca cada dato por su ETIQUETA, nunca por posición fija. Las etiquetas aceptadas están en ETIQUETAS
  (sin tildes ni mayúsculas: la comparación se hace normalizada). Si un modelo usa otra etiqueta, se añade aquí.
- Solo extrae el escenario de PRECIO DE MERCADO. Las columnas o bloques de precio mínimo se ignoran.
- Sin --escribir solo informa: lo encontrado, lo que no ha encontrado y las diferencias con el project.json actual.
- Con --escribir actualiza en project.json únicamente las cifras (proyecto, plan_pagos, cronograma, packs) y solo si
  TODAS las validaciones cuadran. Si algo no cuadra, para y muestra la diferencia exacta (código de salida 2).
- Los textos, rangos (los fija dirección), renders y zona no se tocan nunca.
Requiere openpyxl (pip install --user openpyxl).
"""
import argparse
import datetime as dt
import json
import os
import re
import sys
import unicodedata

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build as B  # noqa: E402

# ------------------------------------------------------------------ etiquetas (normalizadas)
ETIQUETAS = {
    # Hoja «Detalle inversión» (proyecto completo)
    'coste_iva': ['TOTAL INVESTMENT', 'TOTAL INVESTMENT IVA INCLUIDO', 'TOTAL INVERSION', 'TOTAL INVERSION CON IVA', 'INVERSION TOTAL CON IVA'],
    'coste_sin_iva': ['TOTAL INVESTMENT SIN IVA', 'TOTAL INVESTMENT (SIN IVA)', 'TOTAL INVERSION SIN IVA', 'INVERSION TOTAL SIN IVA'],
    'ventas': ['VENTAS', 'TOTAL VENTAS', 'VENTAS TOTALES', 'INGRESOS POR VENTAS', 'TOTAL SALES', 'GDV'],
    # Hoja TIR
    'tir': ['TIR', 'TIR PROYECTO', 'TIR DEL PROYECTO', 'IRR', 'PROJECT IRR'],
    # Plan de pagos
    'aportas': ['APORTAS', 'TOTAL APORTADO', 'TOTAL APORTACIONES', 'APORTACION TOTAL'],
    'recibes': ['RECIBES', 'TOTAL RECIBIDO', 'TOTAL COBROS', 'TOTAL A RECIBIR'],
    'ganancia_neta': ['GANANCIA NETA', 'BENEFICIO NETO DE CAJA', 'RESULTADO NETO DE CAJA'],
    # Cashflow mensual (filas)
    'cf_costes': ['TOTAL COSTES', 'COSTES', 'TOTAL PAGOS', 'PAGOS', 'TOTAL OUTFLOWS'],
    'cf_cobros': ['TOTAL COBROS', 'COBROS', 'TOTAL INGRESOS', 'INGRESOS', 'TOTAL INFLOWS'],
    # Bloque de pack (referencia Montserrat)
    'p_gla': ['GLA'],
    'p_year': ['YEAR BUILT'],
    'p_purchase': ['PURCHASE PRICE'],
    'p_adq': ['ADQUISICION'],
    'p_hard': ['HARD COST'],
    'p_soft': ['SOFT COST'],
    'p_hardsoft': ['HARD&SOFT COST', 'HARD & SOFT COST'],
    'p_fees': ['FEES&FINANCIACION', 'FEES & FINANCIACION'],
    'p_total': ['TOTAL INVESTMENT'],
    'p_total_sin_iva': ['TOTAL INVESTMENT SIN IVA', 'TOTAL INVESTMENT (SIN IVA)'],
    'p_ventas': ['VENTAS', 'TOTAL VENTAS', 'GDV'],
    'p_tir': ['TIR', 'IRR'],
}
# Cabeceras de la tabla de viviendas de cada pack
CAB_VIVIENDAS = {
    'portal': ['PORTAL'], 'planta': ['PLANTA'], 'num': ['N', 'NO', 'NUM', 'NUMERO', 'VIVIENDA', 'PUERTA'],
    'm2': ['M2', 'M2 CONSTRUIDOS', 'M2 CONSTRUIDOS CON COMUNES', 'SUPERFICIE'], 'tipo': ['TIPOLOGIA', 'HAB', 'DORMITORIOS'],
    'orient': ['ORIENTACION'], 'precio': ['PRECIO MERCADO', 'PRECIO DE MERCADO', 'PVP MERCADO'],
}
HOJAS = {
    'detalle': r'detalle\s*(de\s*)?inversi',
    'tir': r'^\s*(tir|irr)\b',
    'cashflow': r'cash\s*-?\s*flow|flujo',
    'pack': r'pack',
}
MERCADO = re.compile(r'MERCADO|MARKET')
MINIMO = re.compile(r'MINIMO|MINIMUM')


def norm(v):
    s = unicodedata.normalize('NFKD', str(v)).encode('ascii', 'ignore').decode()
    s = re.sub(r'[\s:]+', ' ', s.upper()).strip()
    return s


def num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


class Hoja:
    def __init__(self, ws):
        self.ws = ws
        self.celdas = [[c.value for c in fila] for fila in ws.iter_rows()]
        self.col_mercado = self._columnas(MERCADO)
        self.col_minimo = self._columnas(MINIMO)

    def _columnas(self, rx):
        cols = set()
        for fila in self.celdas:
            for j, v in enumerate(fila):
                if isinstance(v, str) and rx.search(norm(v)):
                    cols.add(j)
        return cols

    def buscar(self, etiquetas, exacta=True):
        """Celdas (fila, col) cuyo texto normalizado coincide con alguna etiqueta."""
        ets = [norm(e) for e in etiquetas]
        out = []
        for i, fila in enumerate(self.celdas):
            for j, v in enumerate(fila):
                if isinstance(v, str):
                    n = norm(v)
                    if (n in ets) if exacta else any(n.startswith(e) for e in ets):
                        out.append((i, j))
        return out

    def valor(self, etiquetas):
        """Primer número a la derecha de la etiqueta, en la columna de precio de mercado si la hoja distingue escenarios."""
        for i, j in self.buscar(etiquetas):
            fila = self.celdas[i]
            cands = [(k, fila[k]) for k in range(j + 1, len(fila)) if num(fila[k])]
            if self.col_minimo:
                cands = [(k, v) for k, v in cands if k not in self.col_minimo]
            if self.col_mercado:
                pref = [(k, v) for k, v in cands if k in self.col_mercado]
                cands = pref or cands
            if cands:
                return cands[0][1], f'{self.ws.title}!{celda(i, cands[0][0])} («{fila[j]}»)'
        return None, None


def celda(i, j):
    col = ''
    j += 1
    while j:
        j, r = divmod(j - 1, 26)
        col = chr(65 + r) + col
    return f'{col}{i + 1}'


def hojas(wb, clave):
    return [wb[n] for n in wb.sheetnames if re.search(HOJAS[clave], n, re.I)]


def trimestres(fechas, valores):
    """Agrupa meses por trimestre natural; si el último trimestre está incompleto, sus meses se suman al anterior."""
    grupos = []
    for f, v in zip(fechas, valores):
        q = f'Q{(f.month - 1) // 3 + 1} {f.year}'
        if grupos and grupos[-1][0] == q:
            grupos[-1][1].append(v)
        else:
            grupos.append([q, [v]])
    if len(grupos) > 1 and len(grupos[-1][1]) < 3:
        q, vs = grupos.pop()
        grupos[-1][1].extend(vs)
    return [q for q, _ in grupos], [round(sum(vs), 2) for _, vs in grupos]


def tir_mensual(flujos):
    """TIR anual equivalente desde flujos mensuales (bisección)."""
    def vpn(r):
        return sum(f / (1 + r) ** k for k, f in enumerate(flujos))
    lo, hi = -0.99, 1.0
    if vpn(lo) * vpn(hi) > 0:
        return None
    for _ in range(200):
        mid = (lo + hi) / 2
        if vpn(lo) * vpn(mid) <= 0:
            hi = mid
        else:
            lo = mid
    return (1 + (lo + hi) / 2) ** 12 - 1


def fila_meses(h):
    """Fila con al menos 6 fechas consecutivas (cabecera de meses del cashflow)."""
    for i, fila in enumerate(h.celdas):
        idx = [j for j, v in enumerate(fila) if isinstance(v, (dt.date, dt.datetime))]
        if len(idx) >= 6:
            return i, idx
    return None, None


def extraer(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=False)
    res, fuentes, falta = {}, {}, []

    def tomar(clave, hs, destino=None):
        for ws in hs:
            v, src = Hoja(ws).valor(ETIQUETAS[clave])
            if v is not None:
                (destino if destino is not None else res)[clave] = v
                fuentes[clave] = src
                return v
        falta.append(clave)
        return None

    det = hojas(wb, 'detalle')
    if not det:
        falta.append('hoja «Detalle inversión»')
    for k in ('coste_iva', 'coste_sin_iva', 'ventas'):
        tomar(k, det)
    tomar('tir', hojas(wb, 'tir') + det)
    for k in ('aportas', 'recibes', 'ganancia_neta'):
        tomar(k, list(wb.worksheets))

    # Cashflow mensual
    cf = hojas(wb, 'cashflow')
    res['cronograma'] = None
    for ws in cf:
        h = Hoja(ws)
        i, idx = fila_meses(h)
        if i is None:
            continue
        fechas = [h.celdas[i][j] for j in idx]
        filas = {}
        for k in ('cf_costes', 'cf_cobros'):
            pos = h.buscar(ETIQUETAS[k])
            if pos:
                r = pos[0][0]
                filas[k] = [h.celdas[r][j] if num(h.celdas[r][j]) else 0 for j in idx]
                fuentes[k] = f'{ws.title}!fila {r + 1} («{h.celdas[r][pos[0][1]]}»)'
        if len(filas) == 2:
            costes = [-abs(x) if x > 0 and all(y >= 0 for y in filas['cf_costes']) else x for x in filas['cf_costes']]
            q, c = trimestres(fechas, costes)
            _, s = trimestres(fechas, filas['cf_cobros'])
            res['cronograma'] = {'trimestres': q, 'costes': c, 'cobros': s}
            res['_flujos_mensuales'] = [a + b for a, b in zip(costes, filas['cf_cobros'])]
            break
    if not res['cronograma']:
        falta.append('cashflow mensual (fila de meses + filas de costes y cobros)')

    # Packs
    packs = []
    for ws in hojas(wb, 'pack'):
        m = re.search(r'(\d+)', ws.title)
        if not m:
            continue
        h = Hoja(ws)
        p = {'id': int(m.group(1))}
        for k in ('p_gla', 'p_total', 'p_total_sin_iva', 'p_ventas', 'p_tir'):
            v, src = h.valor(ETIQUETAS[k])
            p[k] = v
            if v is None:
                falta.append(f'{ws.title}: {ETIQUETAS[k][0]}')
            else:
                fuentes[f'pack{p["id"]}.{k}'] = src
        p['u'] = viviendas(h, falta)
        packs.append(p)
    res['_packs'] = sorted(packs, key=lambda p: p['id'])
    return res, fuentes, falta


def viviendas(h, falta):
    """Tabla de viviendas del pack, localizada por cabeceras."""
    cab = {k: [norm(x) for x in v] for k, v in CAB_VIVIENDAS.items()}
    for i, fila in enumerate(h.celdas):
        cols = {}
        for j, v in enumerate(fila):
            if isinstance(v, str):
                n = norm(v)
                for k, vs in cab.items():
                    if n in vs and k not in cols:
                        cols[k] = j
        if {'m2', 'tipo', 'precio'} <= set(cols):
            out = []
            for fila2 in h.celdas[i + 1:]:
                if not num(fila2[cols['m2']]) or not num(fila2[cols['precio']]):
                    if out:
                        break
                    continue
                tipo = fila2[cols['tipo']]
                tn = 0 if isinstance(tipo, str) and 'ESTUDIO' in norm(tipo) else int(re.sub(r'\D', '', str(tipo)) or 0)
                g = lambda k: fila2[cols[k]] if k in cols else ''
                out.append([str(g('portal')), str(g('planta')), str(g('num')), round(float(fila2[cols['m2']]), 1), tn,
                            str(g('orient')), int(round(fila2[cols['precio']]))])
            return out
    falta.append(f'{h.ws.title}: tabla de viviendas (cabeceras m², tipología y precio de mercado)')
    return []


def componer(res, falta):
    """Pasa lo extraído al formato del contrato de project.json y ejecuta las validaciones de §4."""
    errores = []
    out = {}
    if all(res.get(k) is not None for k in ('coste_iva', 'coste_sin_iva', 'ventas', 'tir')):
        tir = res['tir'] / 100 if res['tir'] > 1 else res['tir']
        out['proyecto'] = {'coste_iva': round(res['coste_iva']), 'coste_sin_iva': round(res['coste_sin_iva']),
                           'ventas': round(res['ventas']), 'tir': round(tir, 4)}
    if all(res.get(k) is not None for k in ('aportas', 'recibes', 'ganancia_neta')):
        out['plan_pagos'] = {k: round(res[k]) for k in ('aportas', 'recibes', 'ganancia_neta')}
    if res.get('cronograma'):
        out['cronograma'] = res['cronograma']
    pr = out.get('proyecto')
    packs = []
    for p in res.get('_packs', []):
        us = p['u']
        if not us or p['p_total'] is None or p['p_ventas'] is None or p['p_gla'] is None:
            errores.append(f'Pack {p["id"]}: le faltan viviendas, m² o precio (no se puede publicar)')
            continue
        cs = p['p_total_sin_iva']
        mixd = {}
        for u in us:
            mixd[str(u[4])] = mixd.get(str(u[4]), 0) + 1
        q = {'id': p['id'], 'n': len(us), 'mix': mixd, 'm2': round(p['p_gla']), 'c': round(p['p_total']), 'v': round(p['p_ventas'])}
        if cs is not None:
            q['b'] = round(p['p_ventas'] - cs)
            q['r'] = round((p['p_ventas'] - cs) / cs, 4)
        else:
            errores.append(f'Pack {p["id"]}: falta TOTAL INVESTMENT sin IVA (para beneficio y yield)')
        if p['p_tir'] is not None:
            q['t'] = round(p['p_tir'] / 100 if p['p_tir'] > 1 else p['p_tir'], 4)
        if pr:
            q['pesos'] = [round(p['p_total'] / pr['coste_iva'], 8), round(p['p_ventas'] / pr['ventas'], 8)]
        q['u'] = us
        packs.append(q)
    if packs:
        out['packs'] = packs
        if pr:
            tol = len(packs)
            for k, ref, n in (('c', pr['coste_iva'], 'inversiones'), ('v', pr['ventas'], 'ventas')):
                dif = sum(x[k] for x in packs) - ref
                if abs(dif) > tol:
                    errores.append(f'La suma de {n} de los packs ({B.g(sum(x[k] for x in packs))} €) no coincide con el proyecto ({B.g(ref)} €): diferencia {dif:+} €')
    if res.get('_flujos_mensuales') and pr:
        t = tir_mensual(res['_flujos_mensuales'])
        if t is None or abs(t - pr['tir']) > 0.001:
            tir_hoja = pr['tir']
            calc = 'no calculable' if t is None else ('%.2f %%' % (t * 100)).replace('.', ',')
            dif_pp = '—' if t is None else ('%+.2f pp' % ((t - tir_hoja) * 100)).replace('.', ',')
            errores.append(f'TIR calculada desde el cashflow ({calc}) ≠ hoja TIR ({B.coma(tir_hoja * 100, 2)} %): '
                           f'diferencia {dif_pp} (tolerancia ±0,10 pp)')
    if res.get('cronograma') and out.get('plan_pagos'):
        tot = sum(res['cronograma']['costes']) + sum(res['cronograma']['cobros'])
        dif = tot - out['plan_pagos']['ganancia_neta']
        if abs(dif) > 1:
            errores.append(f'Subtotal del cronograma ({tot:,.2f} €) ≠ ganancia neta ({out["plan_pagos"]["ganancia_neta"]} €): diferencia {dif:+.2f} €')
    return out, errores


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('proyecto')
    ap.add_argument('--excel', required=True)
    ap.add_argument('--escribir', action='store_true')
    ap.add_argument('--json', help='guarda aquí lo extraído (fuera del repo)')
    a = ap.parse_args()
    carpeta = B.localizar(a.proyecto)
    pj = os.path.join(carpeta, 'project.json')
    with open(pj, encoding='utf-8') as f:
        data = json.load(f)
    res, fuentes, falta = extraer(a.excel)
    out, errores = componer(res, falta)
    nviv_json = len(B.viviendas(B.Datos(data, data.get('modo') or 'unico')))
    if out.get('packs') and nviv_json and sum(p['n'] for p in out['packs']) != nviv_json:
        errores.append(f'La suma de viviendas de los packs ({sum(p["n"] for p in out["packs"])}) no coincide con el total del proyecto ({nviv_json})')

    print(f'Excel: {a.excel}\nEncontrado:')
    for k, v in sorted(fuentes.items()):
        print(f'  - {k}: {v}')
    if falta:
        print('No encontrado (revisa la etiqueta en el Excel o añádela a ETIQUETAS):')
        for x in falta:
            print(f'  - {x}')
    for k in ('proyecto', 'plan_pagos', 'cronograma'):
        if k in out and any((data.get(k) or {}).get(c) != v for c, v in out[k].items()):
            print(f'Cambia {k}: {json.dumps(data.get(k), ensure_ascii=False)} → {json.dumps(out[k], ensure_ascii=False)}')
    if a.json:
        with open(a.json, 'w', encoding='utf-8') as f:
            json.dump(out, f, ensure_ascii=False, indent=1)
    if errores:
        print('\nPARA: hay descuadres; no se escribe nada ni se publica:')
        for e in errores:
            print(f'  ✗ {e}')
        sys.exit(2)
    if a.escribir:
        if falta:
            print('\nNo se escribe: faltan datos (ver arriba). Pídelos antes de publicar.')
            sys.exit(2)
        for k, v in out.items():
            if isinstance(v, dict) and isinstance(data.get(k), dict):
                data[k].update(v)  # conserva claves adicionales (p. ej. plan_pagos.desglose)
            else:
                data[k] = v
        B.escribir_json(pj, data)
        print(f'\nproject.json actualizado ({", ".join(out)}). Ejecuta build.py y qa.py.')
    else:
        print('\nValidaciones correctas. Usa --escribir para actualizar project.json.')


if __name__ == '__main__':
    main()
